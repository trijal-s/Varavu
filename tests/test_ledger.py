"""End-to-end tests of inbox → drafts → ledger → reports, using a fake provider (no model calls)."""
import pytest
from openpyxl import load_workbook

from backend.app.db import connect
from backend.app.extraction.providers.base import ExtractionProvider
from backend.app.schemas.extraction import BillEntry, DocumentExtraction, LineItem
from backend.app.services.drafts import list_drafts, reject_draft, scan_upload
from backend.app.services.export import export_year
from backend.app.services.inbox import list_pending, register_inbox
from backend.app.services.ledger import confirm_draft, list_transactions
from backend.app.services.reports import monthly_summary, supplier_dues


class FakeProvider(ExtractionProvider):
    name = "fake"

    def __init__(self, doc):
        self.doc = doc

    def extract(self, image_bytes, mime_type, prompt):
        return self.doc


def cattle_feed(rate=580, document_type="cash_bill", balance=None) -> DocumentExtraction:
    rows = [LineItem(description="feed", quantity=10, rate=810, amount=8100),
            LineItem(description="bran", quantity=5, rate=rate, amount=2900)]
    entry = BillEntry(vendor_name="SRI MURUGAN TRADERS", bill_number="2361", date_as_written="22-2-21",
                      line_items=rows, total=11000, balance_due=balance)
    return DocumentExtraction(document_type=document_type, is_readable=True, is_obscured=False, entries=[entry])


@pytest.fixture
def setup(tmp_path):
    """A fresh database and an inbox with one photo."""
    conn = connect(tmp_path / "test.db")
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "bill1.png").write_bytes(b"photo-1")
    return conn, inbox


def scan_one(conn, inbox, doc):
    register_inbox(conn, inbox)
    upload_id = list_pending(conn)[0]["upload_id"]
    return scan_upload(conn, upload_id, FakeProvider(doc))[0]


def test_same_photo_registered_once(setup):
    conn, inbox = setup
    assert len(register_inbox(conn, inbox)) == 1
    assert register_inbox(conn, inbox) == []


def test_scan_creates_draft_with_review_flags(setup):
    conn, inbox = setup
    draft = scan_one(conn, inbox, cattle_feed(rate=680))  # misread rate
    assert "line_items[1].rate" in draft["needs_review"]
    assert list_pending(conn) == []
    assert len(list_drafts(conn)) == 1


def test_confirm_with_correction_goes_to_ledger(setup):
    conn, inbox = setup
    draft = scan_one(conn, inbox, cattle_feed(rate=680))
    result = confirm_draft(conn, draft["draft_id"], "purchase", "cattle feed", {"line_items[1].rate": 580})
    assert result["warnings"] == []
    assert result["date"] == "2021-02-22"
    [row] = list_transactions(conn)
    assert row["total"] == 11000 and row["category"] == "cattle feed"


def test_vendor_category_is_remembered(setup):
    conn, inbox = setup
    first = scan_one(conn, inbox, cattle_feed())
    confirm_draft(conn, first["draft_id"], "purchase", "cattle feed")
    (inbox / "bill2.png").write_bytes(b"photo-2")
    second_doc = cattle_feed()
    second_doc.entries[0].bill_number = "2400"
    second = scan_one(conn, inbox, second_doc)
    assert confirm_draft(conn, second["draft_id"], "purchase")["category"] == "cattle feed"


def test_duplicate_bill_refused(setup):
    conn, inbox = setup
    confirm_draft(conn, scan_one(conn, inbox, cattle_feed())["draft_id"], "purchase")
    (inbox / "bill1_again.png").write_bytes(b"photo-1-retaken")
    draft = scan_one(conn, inbox, cattle_feed())
    with pytest.raises(ValueError, match="duplicate"):
        confirm_draft(conn, draft["draft_id"], "purchase")


def test_quotation_cannot_enter_ledger(setup):
    conn, inbox = setup
    draft = scan_one(conn, inbox, cattle_feed(document_type="quotation"))
    assert "warning" in draft
    with pytest.raises(ValueError, match="quotation"):
        confirm_draft(conn, draft["draft_id"], "purchase")
    assert reject_draft(conn, draft["draft_id"], "quotation, not a bill")["status"] == "rejected"


def test_transaction_type_required(setup):
    conn, inbox = setup
    draft = scan_one(conn, inbox, cattle_feed())
    with pytest.raises(ValueError, match="transaction_type"):
        confirm_draft(conn, draft["draft_id"])


def test_reports_and_export(setup, tmp_path):
    conn, inbox = setup
    confirm_draft(conn, scan_one(conn, inbox, cattle_feed(balance=3000))["draft_id"], "purchase", "cattle feed")
    summary = monthly_summary(conn, "2021-02")
    assert summary["by_type"]["purchase"]["amount"] == 11000
    assert supplier_dues(conn) == [{"vendor": "SRI MURUGAN TRADERS", "amount_due": 3000, "bills": 1,
                                    "oldest_bill": "2021-02-22"}]
    result = export_year(conn, "2020-21", tmp_path / "exports")
    sheet = load_workbook(result["file"])["Transactions"]
    assert result["transactions"] == 1 and sheet.cell(row=2, column=7).value == 11000
