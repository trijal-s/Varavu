"""Year-end Excel export for the CA: every transaction, every line item, and monthly totals."""
import os
import re
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def parse_financial_year(text: str) -> int:
    """'2023-24' or '2023' -> 2023 (the year the financial year starts in April)."""
    match = re.fullmatch(r"(20\d{2})(?:-(\d{2}|\d{4}))?", text.strip())
    if not match:
        raise ValueError("Use a financial year like 2023-24.")
    return int(match.group(1))


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
        sheet.column_dimensions[get_column_letter(i)].width = min(width + 2, 45)
    sheet.freeze_panes = "A2"


def export_year(conn: sqlite3.Connection, financial_year: str, folder: Path | None = None) -> dict:
    start = parse_financial_year(financial_year)
    first, last = f"{start}-04-01", f"{start + 1}-03-31"
    label = f"{start}-{(start + 1) % 100:02d}"

    transactions = conn.execute("""SELECT * FROM transactions WHERE bill_date BETWEEN ? AND ?
                                   ORDER BY bill_date, id""", (first, last)).fetchall()
    workbook = Workbook()
    workbook.remove(workbook.active)

    add_sheet(workbook, "Transactions",
              ["ID", "Date", "Type", "Category", "Vendor", "Bill no.", "Total", "Tax", "Paid", "Balance", "Payment"],
              [[t["id"], t["bill_date"], t["transaction_type"], t["category"], t["vendor_name"], t["bill_number"],
                t["total"], t["tax_total"], t["amount_paid"], t["balance_due"], t["payment_mode"]] for t in transactions])

    items = conn.execute("""SELECT t.id, t.bill_date, t.vendor_name, l.description, l.quantity, l.unit, l.rate,
                                   l.amount, l.tax_rate
                            FROM line_items l JOIN transactions t ON t.id = l.transaction_id
                            WHERE t.bill_date BETWEEN ? AND ? ORDER BY t.bill_date, t.id""", (first, last)).fetchall()
    add_sheet(workbook, "Line items",
              ["Transaction", "Date", "Vendor", "Item", "Qty", "Unit", "Rate", "Amount", "GST %"],
              [list(row) for row in items])

    monthly = conn.execute("""SELECT substr(bill_date, 1, 7) AS month, transaction_type, category,
                                     round(sum(total), 2) AS amount, count(*) AS bills
                              FROM transactions WHERE bill_date BETWEEN ? AND ?
                              GROUP BY month, transaction_type, category ORDER BY month""", (first, last)).fetchall()
    add_sheet(workbook, "Monthly totals", ["Month", "Type", "Category", "Amount", "Bills"],
              [list(row) for row in monthly])

    folder = Path(folder or os.getenv("VARAVU_EXPORTS", "data/exports"))
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"varavu_FY{label}.xlsx"
    workbook.save(path)
    return {"file": str(path.resolve()), "financial_year": label, "transactions": len(transactions)}
